#!/usr/bin/env python3
"""WhatsApp-Einpager (Hochformat, Handy) fuer den letzten Doppelkopf-Abend.

Nutzung:  python3 doko_whatsapp_pdf.py <Doppelkopf_Datenbank.xlsx> [Ausgabe.pdf] [TT.MM.JJJJ]
Erzeugt HTML und rendert es per Playwright/Chromium als PDF.
"""
import sys, collections, html, subprocess, os, tempfile, json
import openpyxl

PLAYERS = ['Sebastian', 'Andy', 'Frede', 'Fabian', 'Chris']
COL = {'Sebastian': 4, 'Andy': 6, 'Frede': 8, 'Fabian': 10, 'Chris': 12}  # Punkte-Spalte (1-basiert), Letter = +1
COLORS = {'Sebastian': '#e8874a', 'Andy': '#4a9ee8', 'Frede': '#4caf80', 'Fabian': '#c47ee8', 'Chris': '#e84a4a'}
HOST = {'Basti': 'Sebastian'}
SOLO = {'B': 'Buben', 'D': 'Damen', 'T': 'Trumpf', 'S': 'Stilles', 'F': 'Solo', 'A': 'Solo'}


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    sessions = collections.OrderedDict()
    for r in range(5, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        if not v[0]:
            continue
        sessions.setdefault(v[0].strftime('%d.%m.%Y'), []).append(v)
    return sessions


def pts(x, p):
    v = x[COL[p] - 1]
    if v is None or (v == 0 and str(x[COL[p]] or '') == 'M'):
        return None
    return v


def letter(x, p):
    return str(x[COL[p]] or '').strip()


def solo_player(x):
    t = x[16]
    if not t:
        return None
    c = [p for p in PLAYERS if letter(x, p) == t]
    if len(c) == 1:
        return c[0]
    vals = {p: pts(x, p) for p in PLAYERS if pts(x, p) not in (None, 0)}
    c = [p for p, v in vals.items() if all(abs(v) == 3 * abs(u) for q, u in vals.items() if q != p)]
    return c[0] if len(c) == 1 else None


def f(v):
    return f'+{v}' if v > 0 else str(v)


def build(sessions, date):
    rows = sessions[date]
    ort = rows[0][1]
    present = [p for p in PLAYERS if any(pts(x, p) is not None for x in rows)]
    tot = {p: sum(pts(x, p) or 0 for x in rows) for p in present}
    rank = sorted(present, key=lambda p: -tot[p])
    played = {p: sum(pts(x, p) is not None for x in rows) for p in present}
    won = {p: sum((pts(x, p) or 0) > 0 for x in rows) for p in present}

    # Verlauf
    cum = {p: 0 for p in present}
    hist = {p: [0] for p in present}
    for x in rows:
        for p in present:
            cum[p] += pts(x, p) or 0
            hist[p].append(cum[p])

    # Highlights
    cells = [(pts(x, p), p, i + 1) for i, x in enumerate(rows) for p in present if pts(x, p) is not None]
    best, worst = max(cells), min(cells)
    streak = {}
    for p in present:
        b = c = 0
        for x in rows:
            v = pts(x, p)
            if v is None:
                continue
            c = c + 1 if v > 0 else 0
            b = max(b, c)
        streak[p] = b
    sp = max(present, key=lambda p: streak[p])

    solos = []
    for i, x in enumerate(rows):
        p = solo_player(x)
        if p:
            solos.append((i + 1, p, x[16], pts(x, p)))
    swon = sum(1 for s in solos if s[3] > 0)
    hz = [(i + 1, p) for i, x in enumerate(rows) for p in present if letter(x, p) == 'H']
    ar = [(i + 1, p) for i, x in enumerate(rows) for p in present if letter(x, p) == 'A']
    bock = sum(1 for x in rows if x[17])
    schw = sum(1 for x in rows if x[14])

    # Ewige Tabelle vorher/nachher
    dates = list(sessions)
    idx = dates.index(date)
    def alltime(upto):
        t = {p: 0 for p in PLAYERS}
        for d in dates[:upto]:
            for x in sessions[d]:
                for p in PLAYERS:
                    t[p] += pts(x, p) or 0
        return t
    before, after = alltime(idx), alltime(idx + 1)
    rb = {p: i for i, p in enumerate(sorted(PLAYERS, key=lambda p: -before[p]))}
    ra = sorted(PLAYERS, key=lambda p: -after[p])

    # SVG Verlauf
    W, H = 460, 150
    lo = min(min(v) for v in hist.values()); hi = max(max(v) for v in hist.values())
    X = lambda i: 30 + i * (W - 40) / len(rows)
    Y = lambda v: 8 + (hi - v) * (H - 16) / ((hi - lo) or 1)
    svg = [f'<svg viewBox="0 0 {W} {H}" width="100%"><line x1="30" x2="{W-10}" y1="{Y(0):.1f}" y2="{Y(0):.1f}" stroke="#b9b3a3" stroke-dasharray="3 3"/>']
    for p in present:
        pl = ' '.join(f'{X(i):.1f},{Y(v):.1f}' for i, v in enumerate(hist[p]))
        svg.append(f'<polyline fill="none" stroke="{COLORS[p]}" stroke-width="{3 if p == rank[0] else 2}" stroke-linejoin="round" points="{pl}"/>')
        svg.append(f'<circle cx="{X(len(rows)):.1f}" cy="{Y(hist[p][-1]):.1f}" r="3.5" fill="{COLORS[p]}"/>')
    svg.append(f'<text x="0" y="{Y(hi)+4:.0f}" font-size="10" fill="#8a8471">{f(hi)}</text><text x="0" y="{Y(lo)+4:.0f}" font-size="10" fill="#8a8471">{lo}</text></svg>')

    mx = max(abs(v) for v in tot.values()) or 1
    lines = []
    for i, p in enumerate(rank[1:], 2):
        w = abs(tot[p]) / mx * 46
        bar = (f'<div class="bar" style="left:50%;width:{w}%;background:{COLORS[p]}"></div>' if tot[p] >= 0 else
               f'<div class="bar" style="left:{50-w}%;width:{w}%;background:{COLORS[p]}"></div>')
        lines.append(f'<div class="rk"><span class="pos">{i}.</span><span class="nm" style="color:{COLORS[p]}">{p}</span>'
                     f'<span class="track"><span class="mid"></span>{bar}</span><span class="pt {"neg" if tot[p]<0 else "pos2"}">{f(tot[p])}</span></div>')

    by_player = collections.defaultdict(list)
    for r, p, t, v in solos:
        by_player[p].append((t, v))
    solo_rows = ''.join(
        f'<div class="sl"><b style="color:{COLORS[p]}">{p}</b> '
        + ' '.join(f'<span class="chip {"w" if v>0 else "l"}">{SOLO.get(t, t)} {f(v)}</span>' for t, v in by_player[p]) + '</div>'
        for p in rank if by_player[p])

    def arrow(p):
        d = rb[p] - ra.index(p)
        return '<span class="up">▲' + str(d) + '</span>' if d > 0 else ('<span class="dn">▼' + str(-d) + '</span>' if d < 0 else '<span class="eq">–</span>')
    table = ''.join(
        f'<tr><td>{i+1}.</td><td style="color:{COLORS[p]};font-weight:700">{p}</td><td class="r">{f(after[p])}</td>'
        f'<td class="r small">{f(after[p]-before[p]) if p in present else "–"}</td><td class="c">{arrow(p)}</td></tr>'
        for i, p in enumerate(ra))

    host = HOST.get(ort, ort)
    w0 = rank[0]
    css = """
    @page{margin:0}
    *{box-sizing:border-box}body{margin:0;font-family:'DejaVu Sans',Arial,sans-serif;background:#f6f2e8;color:#23201a}
    .pg{padding:18px 18px 12px}
    .hd{background:#1f3a2c;color:#f6f2e8;border-radius:14px;padding:14px 16px}
    .hd h1{margin:0;font-size:21px;letter-spacing:.3px}.hd .s{opacity:.8;font-size:12.5px;margin-top:3px}
    .win{margin-top:12px;background:#fff;border-radius:14px;padding:12px 14px;border:2px solid var(--c);display:flex;align-items:center;gap:12px}
    .win .cr{font-size:34px}.win .n{font-size:22px;font-weight:800;color:var(--c)}.win .p{margin-left:auto;font-size:30px;font-weight:800;color:#2e7d4f}
    .win .sub{font-size:11px;color:#6d6758}
    .card{margin-top:10px;background:#fff;border-radius:14px;padding:10px 14px}
    .card h2{margin:0 0 6px;font-size:13px;text-transform:uppercase;letter-spacing:.8px;color:#8a8471}
    .rk{display:flex;align-items:center;gap:8px;margin:5px 0;font-size:14px}.pos{width:18px;color:#8a8471}.nm{width:78px;font-weight:700}
    .track{position:relative;flex:1;height:12px}.mid{position:absolute;left:50%;top:-2px;bottom:-2px;border-left:1px solid #cfc8b6}
    .bar{position:absolute;top:0;height:12px;border-radius:6px}.pt{width:48px;text-align:right;font-weight:800}.neg{color:#c0392b}.pos2{color:#2e7d4f}
    .grid{display:grid;grid-template-columns:1fr 1fr;gap:8px;margin-top:10px}
    .t{background:#fff;border-radius:12px;padding:8px 10px}.t .lb{font-size:10.5px;color:#8a8471;text-transform:uppercase;letter-spacing:.5px}
    .t .v{font-size:19px;font-weight:800;margin-top:2px}.t .d{font-size:11.5px;color:#4b463b}
    .sl{font-size:12.5px;margin:4px 0}.chip{display:inline-block;border-radius:8px;padding:1px 6px;font-size:11.5px;margin:1px 2px;font-weight:700}
    .w{background:#dff1e5;color:#1f6b40}.l{background:#f8dfdb;color:#a0301f}
    .ev{display:flex;justify-content:space-between;font-size:12.5px;margin-top:6px;color:#4b463b}
    table{width:100%;border-collapse:collapse;font-size:13.5px}td{padding:4px 2px;border-bottom:1px solid #eee7d6}
    .r{text-align:right;font-weight:700}.small{font-size:11.5px;color:#6d6758;font-weight:400}.c{text-align:center;width:34px}
    .up{color:#2e7d4f;font-weight:700}.dn{color:#c0392b;font-weight:700}.eq{color:#b9b3a3}
    .ft{text-align:center;font-size:11px;color:#8a8471;margin-top:8px}
    """
    out = f"""<!doctype html><meta charset="utf-8"><style>{css}</style><div class="pg">
<div class="hd"><h1>🃏 Doko-Abend {date}</h1><div class="s">bei {html.escape(host)} · {len(rows)} Runden · {len(present)} Spieler</div></div>
<div class="win" style="--c:{COLORS[w0]}"><div class="cr">👑</div><div><div class="n">{w0}</div><div class="sub">{won[w0]} von {played[w0]} Runden gewonnen</div></div><div class="p">{f(tot[w0])}</div></div>
<div class="card"><h2>Abendwertung</h2>{''.join(lines)}</div>
<div class="card"><h2>Punkteverlauf</h2>{''.join(svg)}</div>
<div class="grid">
 <div class="t"><div class="lb">🚀 Beste Runde</div><div class="v" style="color:{COLORS[best[1]]}">{f(best[0])}</div><div class="d">{best[1]} · Runde {best[2]}</div></div>
 <div class="t"><div class="lb">💥 Bitterste Runde</div><div class="v" style="color:{COLORS[worst[1]]}">{worst[0]}</div><div class="d">{worst[1]} · Runde {worst[2]}</div></div>
 <div class="t"><div class="lb">🔥 Längste Siegesserie</div><div class="v" style="color:{COLORS[sp]}">{streak[sp]} Runden</div><div class="d">{sp}</div></div>
 <div class="t"><div class="lb">🎭 Solo-Bilanz</div><div class="v">{swon} : {len(solos)-swon}</div><div class="d">{len(solos)} Solos gewonnen : verloren</div></div>
</div>
<div class="card"><h2>Solos</h2>{solo_rows or '<div class="sl">keine</div>'}
<div class="ev"><span>🔔 {bock} Bock</span><span>💍 {len(hz)} Hochzeit</span><span>🤲 {len(ar)} Armut</span><span>🐷 {schw} Schwein</span></div></div>
<div class="card"><h2>Ewige Tabelle</h2><table>{table}</table></div>
<div class="ft">Alle Abende & Runden: sebaredo.github.io/DokoRunde</div>
</div>"""
    return out


def main():
    db = sys.argv[1]
    sessions = load(db)
    date = sys.argv[3] if len(sys.argv) > 3 else list(sessions)[-1]
    out = sys.argv[2] if len(sys.argv) > 2 else f'Doko-Abend {date}.pdf'
    h = build(sessions, date)
    tmp = tempfile.mkdtemp()
    hp = os.path.join(tmp, 'wa.html')
    open(hp, 'w').write(h)
    js = f"""const pw=require(process.env.PW_MODULE||'playwright');(async()=>{{
const b=await pw.chromium.launch(process.env.PW_CHROMIUM?{{executablePath:process.env.PW_CHROMIUM}}:{{}});
const p=await b.newPage();await p.goto('file://'+{json.dumps(hp)});
await p.setViewportSize({{width:454,height:800}});const h=await p.evaluate(()=>document.querySelector('.pg').scrollHeight);
await p.pdf({{path:{json.dumps(os.path.abspath(out))},width:'454px',height:(h+2)+'px',printBackground:true,pageRanges:'1'}});
await p.setViewportSize({{width:454,height:805}});await p.screenshot({{path:{json.dumps(os.path.abspath(out)[:-4] + '.png')},fullPage:true}});
await b.close();}})();"""
    jp = os.path.join(tmp, 'r.js')
    open(jp, 'w').write(js)
    subprocess.run(['node', jp], check=True)
    print(out)


if __name__ == '__main__':
    main()
