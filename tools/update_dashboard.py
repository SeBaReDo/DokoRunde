#!/usr/bin/env python3
"""Aktualisiert index.html (Dashboard) komplett aus Doppelkopf_Datenbank.xlsx und prueft danach Dashboard <-> DB.

Nutzung:  python3 tools/update_dashboard.py <Doppelkopf_Datenbank.xlsx> [index.html]
Neu erzeugt: DATA.absences, DATA.sessions, cumulative, wins, final, ABSENCES_MAP, gamesPlayed, streaks,
solosByPlayer, solosLostByPlayer und ROUNDS (Einzelrunden; Event-Codes ab 2026 aus der DB, Altdaten bleiben).
"""
import sys, re, json, collections
import openpyxl

PLAYERS = ['Sebastian', 'Andy', 'Frede', 'Fabian', 'Chris']
COL = {'Sebastian': 4, 'Andy': 6, 'Frede': 8, 'Fabian': 10, 'Chris': 12}
SOLO_CODE = {'B': 'Bs', 'D': 'D', 'T': 'T', 'S': 'St'}
SOLO_E = {'F', 'A', 'D', 'T', 'Bs', 'St'}


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    s = {}
    for r in range(5, ws.max_row + 1):
        v = [c.value for c in ws[r]]
        if v[0]:
            s.setdefault(v[0], []).append(v)
    # chronologisch (die DB-Zeilen sind nicht durchgehend sortiert)
    return collections.OrderedDict((d.strftime('%d.%m.%Y'), s[d]) for d in sorted(s))


def pts(x, p):
    v = x[COL[p] - 1]
    return None if v is None or (v == 0 and str(x[COL[p]] or '') == 'M') else v


def solo_player(x):
    t = x[16]
    if not t:
        return None
    c = [p for p in PLAYERS if str(x[COL[p]] or '').strip() == t]
    if len(c) == 1:
        return c[0]
    vals = {p: pts(x, p) for p in PLAYERS if pts(x, p) not in (None, 0)}
    c = [p for p, v in vals.items() if all(abs(v) == 3 * abs(u) for q, u in vals.items() if q != p)]
    return c[0] if len(c) == 1 else None


def stats(sessions):
    st = {}
    for d, rows in sessions.items():
        present = [p for p in PLAYERS if any(pts(x, p) is not None for x in rows)]
        sc = {p: sum(pts(x, p) or 0 for x in rows) for p in PLAYERS}
        st[d] = dict(ort=rows[0][1], rounds=len(rows), scores=sc, present=present,
                     absent=[p for p in PLAYERS if p not in present],
                     winner=max(present, key=lambda p: sc[p]),
                     hochzeit=sum(bool(x[13]) for x in rows), schwein=sum(bool(x[14]) for x in rows),
                     solo=sum(bool(x[16]) for x in rows), bock=sum(bool(x[17]) for x in rows))
    return st


def player_stats(sessions):
    allrows = [x for rows in sessions.values() for x in rows]
    gp = {p: sum(pts(x, p) is not None for x in allrows) for p in PLAYERS}
    streaks = {}
    for p in PLAYERS:
        bw = bl = 0
        for rows in sessions.values():
            cw = cl = 0
            for x in rows:
                v = pts(x, p)
                if v is None:
                    continue
                cw, cl = (cw + 1, 0) if v > 0 else ((0, cl + 1) if v < 0 else (0, 0))
                bw, bl = max(bw, cw), max(bl, cl)
        streaks[p] = (bw, bl)
    so, sl = collections.Counter(), collections.Counter()
    for x in allrows:
        p = solo_player(x)
        if p:
            so[p] += 1
            sl[p] += pts(x, p) < 0
    return gp, streaks, so, sl


def round_entries(rows):
    out = []
    for i, x in enumerate(rows, 1):
        r = {'r': i}
        for p in PLAYERS:
            if pts(x, p) is not None:
                r[p] = pts(x, p)
        e = []
        if x[16]:
            e.append(SOLO_CODE.get(x[16], 'F'))
        if x[13]: e.append('H')
        if x[15]: e.append('Ar')
        if x[14]: e.append('S')
        if x[17]: e.append('B')
        if e:
            r['e'] = e
        out.append(r)
    return out


def objstr(d):
    return '{' + ','.join(f'{k}:{v}' for k, v in d.items()) + '}'


def update(html, sessions):
    st = stats(sessions)
    dates = list(st)
    # absences
    absn = [(d, s['absent']) for d, s in st.items() if s['absent']]
    html = re.sub(r'absences: \{.*?\n  \},', 'absences: {\n' + ',\n'.join(f'    "{d}": {json.dumps(a)}' for d, a in absn) + '\n  },', html, flags=re.S)
    html = re.sub(r"const ABSENCES_MAP = \{.*?\n\};", 'const ABSENCES_MAP = {\n' + '\n'.join(f"  '{d}': {json.dumps(a)}," for d, a in absn) + '\n};', html, flags=re.S)
    # sessions
    lines = []
    for d, s in st.items():
        ort = json.dumps(s['ort']) + ','
        lines.append(f'    {{ date:"{d}", ort:{ort:<9}rounds:{s["rounds"]}, scores:{objstr(s["scores"])}, winner:"{s["winner"]}", '
                     f'hochzeit:{s["hochzeit"]}, schwein:{s["schwein"]}, solo:{s["solo"]}, bock:{s["bock"]} }}')
    html = re.sub(r'sessions: \[.*?\n  \],', 'sessions: [\n' + ',\n'.join(lines) + '\n  ],', html, flags=re.S)
    # cumulative / wins / final
    cum = {p: [] for p in PLAYERS}
    run = {p: 0 for p in PLAYERS}
    for d in dates:
        for p in PLAYERS:
            run[p] += st[d]['scores'][p]
            cum[p].append(run[p])
    block = '  cumulative: {\n    labels:    ' + json.dumps(dates).replace(', ', ',') + ',\n' + ',\n'.join(
        f'    {p + ":":<11}[{",".join(map(str, cum[p]))}]' for p in PLAYERS) + '\n  },'
    html = re.sub(r'  cumulative: \{.*?\n  \},', block, html, flags=re.S)
    wins = collections.Counter(s['winner'] for s in st.values())
    html = re.sub(r'wins:  \{[^}]*\}', 'wins:  { ' + ', '.join(f'{p}:{wins[p]}' for p in PLAYERS) + ' }', html)
    html = re.sub(r'final: \{[^}]*\}', 'final: { ' + ', '.join(f'{p}:{run[p]}' for p in PLAYERS) + ' }', html)
    # Spielerstatistiken
    gp, streaks, so, sl = player_stats(sessions)
    html = re.sub(r'const gamesPlayed = \{[^}]*\};', 'const gamesPlayed = { ' + ', '.join(f'{p}:{gp[p]}' for p in PLAYERS) + ' };', html)
    html = re.sub(r'const streaks = \{.*?\n\};', 'const streaks = {\n' + ',\n'.join(
        f'  {p + ":":<10} {{ maxWin: {str(streaks[p][0]) + ",":<3} maxLoss: {streaks[p][1]:<2} }}' for p in PLAYERS) + '\n};', html, flags=re.S)
    html = re.sub(r'const solosByPlayer = \{[^}]*\};', 'const solosByPlayer = { ' + ', '.join(f'{p}:{so[p]}' for p in PLAYERS) + ' };', html)
    html = re.sub(r'const solosLostByPlayer = \{[^}]*\};', 'const solosLostByPlayer = { ' + ', '.join(f'{p}:{sl[p]}' for p in PLAYERS) + ' };', html)
    # ROUNDS
    i = html.index('const ROUNDS = ')
    j = html.index('\n', i)
    tail = html[i + 15:j]
    semi = ';' if tail.rstrip().endswith(';') else ''
    old = {r['date']: r for r in json.loads(tail.rstrip().rstrip(';'))}
    R = []
    for d, rows in sessions.items():
        if d in old and not d.endswith(('2026', '2027', '2028', '2029')):
            R.append(old[d])
        else:
            R.append({'date': d, 'rounds': round_entries(rows)})
    html = html[:i + 15] + json.dumps(R, ensure_ascii=False, separators=(',', ':')) + semi + html[j:]
    return html


def verify(html, sessions):
    """Abgleich Dashboard <-> DB. Gibt Liste von Abweichungen zurueck."""
    st = stats(sessions)
    errs = []
    sess = re.findall(r'\{ date:"([^"]+)", ort:"([^"]+)",\s*rounds:(\d+), scores:\{([^}]*)\}, winner:"(\w+)", hochzeit:(\d+), schwein:(\d+), solo:(\d+), bock:(\d+) \}', html)
    if [s[0] for s in sess] != list(st):
        errs.append('Abend-Liste weicht ab')
    for d, ort, n, sc, w, h, sw, so, b in sess:
        s = st.get(d)
        if not s:
            continue
        scd = {k: int(v) for k, v in (kv.split(':') for kv in sc.split(','))}
        for k, v in [('ort', ort == s['ort']), ('rounds', int(n) == s['rounds']), ('scores', scd == s['scores']), ('winner', w == s['winner']),
                     ('hochzeit', int(h) == s['hochzeit']), ('schwein', int(sw) == s['schwein']), ('solo', int(so) == s['solo']), ('bock', int(b) == s['bock'])]:
            if not v:
                errs.append(f'{d}: {k}')
    i = html.index('const ROUNDS = ')
    R = json.loads(html[i + 15:html.index('\n', i)].rstrip().rstrip(';'))
    for e in R:
        rows = sessions[e['date']]
        if len(e['rounds']) != len(rows):
            errs.append(f"{e['date']}: Rundenzahl ROUNDS")
        for r, x in zip(e['rounds'], rows):
            for p in PLAYERS:
                if r.get(p) != pts(x, p):
                    errs.append(f"{e['date']} R{r['r']} {p}: {r.get(p)} != {pts(x, p)}")
        if sum(any(c in SOLO_E for c in r.get('e', [])) for r in e['rounds']) != st[e['date']]['solo']:
            errs.append(f"{e['date']}: Solo-Codes in ROUNDS")
        if e['date'].endswith('2026') and sum(r.get('e', []).count('B') for r in e['rounds']) != st[e['date']]['bock']:
            errs.append(f"{e['date']}: Bock-Codes in ROUNDS")
    a1 = re.search(r'absences: \{(.*?)\n  \}', html, re.S).group(1)
    a2 = re.search(r'const ABSENCES_MAP = \{(.*?)\n\};', html, re.S).group(1)
    exp = {d: s['absent'] for d, s in st.items() if s['absent']}
    if {d: json.loads(a) for d, a in re.findall(r'"([\d.]+)": (\[[^\]]*\])', a1)} != exp:
        errs.append('DATA.absences')
    if {d: json.loads(a) for d, a in re.findall(r"'([\d.]+)': (\[[^\]]*\])", a2)} != exp:
        errs.append('ABSENCES_MAP')
    return errs, st


if __name__ == '__main__':
    db = sys.argv[1]
    path = sys.argv[2] if len(sys.argv) > 2 else 'index.html'
    sessions = load(db)
    html = update(open(path, encoding='utf-8').read(), sessions)
    open(path, 'w', encoding='utf-8').write(html)
    errs, st = verify(html, sessions)
    print(f'{"Datum":<11} {"Ort":<7} Rd  Sieger     H  Sw So Bo  Abwesend')
    for d, s in st.items():
        print(f'{d:<11} {s["ort"]:<7} {s["rounds"]:>2}  {s["winner"]:<9} {s["hochzeit"]:>2} {s["schwein"]:>2} {s["solo"]:>2} {s["bock"]:>2}  {",".join(s["absent"])}')
    print('Abgleich Dashboard <-> DB:', 'OK, keine Abweichungen' if not errs else errs)
    sys.exit(1 if errs else 0)
